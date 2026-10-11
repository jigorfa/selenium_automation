"""Mescla a resposta do NotebookLM (Markdown) no modelo .docx do relatório.

Como funciona: o modelo é aberto como base (cabeçalho com logo, rodapé, fonte
embutida e definições de lista são mantidos), o corpo de exemplo é apagado e a
resposta é escrita no lugar, reproduzindo a formatação do modelo:

    - "Relatório de Edital" e "Resumo edital": título centralizado, 14 pt
    - "Pontos de atenção": negrito com marca-texto vermelho
    - texto comum: justificado; **negrito** e *itálico* viram formatação do Word
    - listas com marcadores e listas numeradas (1. > a. > i. > 1.) usam a
      numeração do próprio modelo
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from config import MODELO_DOCX, MODELOS_DOCX_POR_PERFIL, REMOVER_CITACOES

# Linhas que, sozinhas, são títulos centralizados (comparação sem acento de caixa/negrito)
TITULOS_CENTRALIZADOS = {"relatório de edital", "resumo edital"}

# Linhas que ganham negrito + marca-texto vermelho, como no modelo
DESTAQUE_VERMELHO = {"pontos de atenção"}

# Subitem com marcador dentro de item numerado vira a/b/c, como no modelo
SUBITENS_COMO_NUMERADOS = True

NIVEL_MAXIMO = 8  # as listas do modelo têm 9 níveis (0 a 8)

RE_TITULO = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
RE_MARCADOR = re.compile(r"^(\s*)[-*+•]\s+(.*)$")
RE_ORDENADO = re.compile(r"^(\s*)(\d{1,3})[.)]\s+(.*)$")
RE_LETRA = re.compile(r"^(\s+)(?:[a-z]|[ivxlc]{1,4})[.)]\s+(.*)$", re.IGNORECASE)
RE_REGUA = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
RE_LINHA_TABELA = re.compile(r"^\s*\|.*\|\s*$")
RE_SEPARADOR_TABELA = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")

RE_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
RE_CITACAO = re.compile(r"\s?\[\d+(?:\s*[,\u2013-]\s*\d+)*\]")
RE_INLINE = re.compile(r"\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(?!\s)(.+?)(?<!\s)\*")


# ---------------------------------------------------------------------------
# 1) Markdown -> blocos
# ---------------------------------------------------------------------------
def _nivel(pilha, tipos, recuo):
    """Descobre o nível da lista pela indentação (funciona com 2 ou 4 espaços)."""
    while pilha and recuo < pilha[-1]:
        pilha.pop()
        tipos.pop()
    if not pilha or recuo > pilha[-1]:
        pilha.append(recuo)
        tipos.append(None)
    return len(pilha) - 1


def _como_item(linha, em_lista):
    """Devolve (recuo, ordenado, numero, texto) se a linha for item de lista."""
    m = RE_MARCADOR.match(linha)
    if m:
        return len(m.group(1)), False, None, m.group(2)
    m = RE_ORDENADO.match(linha)
    if m:
        return len(m.group(1)), True, int(m.group(2)), m.group(3)
    m = RE_LETRA.match(linha)
    if m and em_lista:  # "a." / "i." só valem indentados dentro de uma lista
        return len(m.group(1)), True, None, m.group(2)
    return None


def interpretar(markdown):
    linhas = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocos = []
    pilha, tipos = [], []
    i = 0

    while i < len(linhas):
        linha = linhas[i].expandtabs(4).rstrip()
        i += 1

        if not linha.strip() or RE_REGUA.match(linha):
            continue

        linha = re.sub(r"^\s*>\s?", "", linha)  # citação (>) vira texto normal

        # Tabela: linha com | seguida da linha separadora |---|---|
        if RE_LINHA_TABELA.match(linha) and i < len(linhas) and RE_SEPARADOR_TABELA.match(linhas[i]):
            rows = [linha]
            i += 1  # pula o separador
            while i < len(linhas) and RE_LINHA_TABELA.match(linhas[i]):
                rows.append(linhas[i])
                i += 1
            pilha.clear()
            tipos.clear()
            blocos.append({
                "tipo": "tabela",
                "linhas": [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows],
            })
            continue

        m = RE_TITULO.match(linha)
        if m:
            pilha.clear()
            tipos.clear()
            blocos.append({"tipo": "titulo", "nivel": len(m.group(1)), "texto": m.group(2)})
            continue

        item = _como_item(linha, bool(pilha))
        if item:
            recuo, ordenado, numero, texto = item
            nivel = _nivel(pilha, tipos, recuo)
            if (SUBITENS_COMO_NUMERADOS and not ordenado
                    and nivel > 0 and tipos[nivel - 1] == "ord"):
                ordenado = True
            tipos[nivel] = "ord" if ordenado else "marc"
            blocos.append({"tipo": "item", "ordenado": ordenado, "nivel": nivel,
                           "numero": numero, "texto": texto})
            continue

        # Linha indentada logo após um item: continuação do texto do item
        if pilha and blocos and blocos[-1]["tipo"] == "item" and linha.startswith(" "):
            blocos[-1]["texto"] += " " + linha.strip()
            continue

        pilha.clear()
        tipos.clear()
        blocos.append({"tipo": "paragrafo", "texto": linha.strip()})

    return blocos


# ---------------------------------------------------------------------------
# 2) Texto inline (**negrito**, *itálico*)
# ---------------------------------------------------------------------------
def _limpar(texto):
    texto = RE_LINK.sub(lambda m: m.group(1) if m.group(1) == m.group(2)
                        else f"{m.group(1)} ({m.group(2)})", texto)
    if REMOVER_CITACOES:
        texto = RE_CITACAO.sub("", texto)
    texto = re.sub(r"`([^`]*)`", r"\1", texto)
    texto = texto.replace("\\*", "\x00")
    texto = re.sub(r"\\([\\_{}\[\]()#+\-.!|>])", r"\1", texto)
    return texto


def _partes(texto, negrito=False, italico=False):
    saida, pos = [], 0
    for m in RE_INLINE.finditer(texto):
        if m.start() > pos:
            saida.append((texto[pos:m.start()], negrito, italico))
        if m.group(1) is not None:
            saida += _partes(m.group(1), True, True)
        elif m.group(2) is not None:
            saida += _partes(m.group(2), True, italico)
        else:
            saida += _partes(m.group(3), negrito, True)
        pos = m.end()
    if pos < len(texto):
        saida.append((texto[pos:], negrito, italico))
    return saida


def _texto_puro(texto):
    """Texto sem negrito/pontuação final, minúsculo - para reconhecer títulos."""
    t = _limpar(texto).replace("*", "").strip().rstrip(":").strip()
    return t.lower()


# ---------------------------------------------------------------------------
# 3) Escrita no .docx
# ---------------------------------------------------------------------------
class _Numeracao:
    """Reaproveita as listas definidas no modelo (numbering.xml)."""

    def __init__(self, doc):
        self.raiz = doc.part.numbering_part.element
        self.mapa = {
            int(n.get(qn("w:numId"))): int(n.find(qn("w:abstractNumId")).get(qn("w:val")))
            for n in self.raiz.findall(qn("w:num"))
        }
        # No modelo: lista 4 = 1. > a. > i. > 1.  |  lista 2 = marcadores
        num_ordenada = self._escolher("decimal", preferido=4)
        self.abstrato_ordenado = self.mapa[num_ordenada]
        self.marcador = self._escolher("bullet", preferido=2)

    def _formato_nivel0(self, abstrato_id):
        for a in self.raiz.findall(qn("w:abstractNum")):
            if int(a.get(qn("w:abstractNumId"))) == abstrato_id:
                for lvl in a.findall(qn("w:lvl")):
                    if lvl.get(qn("w:ilvl")) == "0":
                        fmt = lvl.find(qn("w:numFmt"))
                        return fmt.get(qn("w:val")) if fmt is not None else None
        return None

    def _escolher(self, formato, preferido):
        if preferido in self.mapa and self._formato_nivel0(self.mapa[preferido]) == formato:
            return preferido
        for num_id, abstrato in self.mapa.items():
            if self._formato_nivel0(abstrato) == formato:
                return num_id
        raise RuntimeError(f"O modelo .docx não tem nenhuma lista do tipo '{formato}'.")

    def nova_ordenada(self, inicio=None):
        """Cria uma nova lista numerada (reinicia a contagem), opcionalmente em `inicio`."""
        novo = max(self.mapa) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(novo))
        abstrato = OxmlElement("w:abstractNumId")
        abstrato.set(qn("w:val"), str(self.abstrato_ordenado))
        num.append(abstrato)
        if inicio is not None:
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            start = OxmlElement("w:startOverride")
            start.set(qn("w:val"), str(inicio))
            override.append(start)
            num.append(override)
        self.raiz.append(num)
        self.mapa[novo] = self.abstrato_ordenado
        return novo


def _fonte_do_modelo(doc):
    for par in doc.paragraphs:
        for run in par.runs:
            if run.font.name:
                return run.font.name
    return None


def _escrever(par, texto, fonte, negrito=None, tamanho=None, vermelho=False):
    for trecho, bold, ital in _partes(_limpar(texto)):
        run = par.add_run(trecho.replace("\x00", "*"))
        if fonte:
            run.font.name = fonte
            rfonts = run._r.get_or_add_rPr().rFonts
            rfonts.set(qn("w:cs"), fonte)
            rfonts.set(qn("w:eastAsia"), fonte)
        if negrito is None:
            run.bold = bold or None
        else:
            run.bold = negrito or None
        run.italic = ital or None
        if tamanho:
            run.font.size = Pt(tamanho)
        if vermelho:
            run.font.highlight_color = WD_COLOR_INDEX.RED


def _numerar(par, num_id, nivel):
    numpr = par._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = nivel
    numpr.get_or_add_numId().val = num_id


def _fonte_na_marca(par, fonte):
    """Fonte da marca do parágrafo: define o tamanho/estilo do símbolo/número da lista."""
    if not fonte:
        return
    rfonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), fonte)
    rpr = OxmlElement("w:rPr")
    rpr.append(rfonts)
    par._p.get_or_add_pPr().append(rpr)


def _titulo(doc, texto, fonte):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _escrever(par, texto, fonte, negrito=False, tamanho=14)


def _tabela(doc, linhas, fonte):
    colunas = max(len(l) for l in linhas)
    tabela = doc.add_table(rows=len(linhas), cols=colunas)

    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "808080")
        bordas.append(b)
    tbl_pr = tabela._tbl.tblPr
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(bordas)
    else:
        tbl_pr.append(bordas)

    for r, linha in enumerate(linhas):
        for c in range(colunas):
            par = tabela.cell(r, c).paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            _escrever(par, linha[c] if c < len(linha) else "", fonte,
                      negrito=True if r == 0 else None)
    doc.add_paragraph()  # espaço depois da tabela


def modelo_para(perfil):
    return MODELOS_DOCX_POR_PERFIL.get(perfil, MODELO_DOCX)


def gerar_relatorio(markdown, destino, modelo=MODELO_DOCX):
    """Escreve `markdown` dentro de uma cópia do `modelo` e salva em `destino`."""
    if not os.path.exists(modelo):
        raise FileNotFoundError(f"Modelo .docx não encontrado: {modelo}")

    doc = Document(modelo)
    fonte = _fonte_do_modelo(doc)
    numeracao = _Numeracao(doc)

    # Apaga o conteúdo de exemplo, mantendo a configuração da página (sectPr),
    # que é quem liga o cabeçalho e o rodapé.
    corpo = doc.element.body
    for filho in list(corpo):
        if filho.tag != qn("w:sectPr"):
            corpo.remove(filho)

    corrida = None  # lista numerada em andamento (id do w:num)

    blocos = interpretar(markdown)

    for posicao, bloco in enumerate(blocos):
        tipo = bloco["tipo"]

        if tipo == "item":
            if bloco["ordenado"]:
                if corrida is None:
                    inicio = bloco["numero"] if bloco["nivel"] == 0 else None
                    corrida = numeracao.nova_ordenada(inicio)
                num_id = corrida
            else:
                if bloco["nivel"] == 0:
                    corrida = None  # lista com marcadores encerra a numerada
                num_id = numeracao.marcador

            par = doc.add_paragraph()
            _numerar(par, num_id, min(bloco["nivel"], NIVEL_MAXIMO))
            # Itens colados entre si; o último da lista mantém o espaço normal depois
            proximo_e_item = (posicao + 1 < len(blocos)
                              and blocos[posicao + 1]["tipo"] == "item")
            if proximo_e_item:
                par.paragraph_format.space_after = Pt(0)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _escrever(par, bloco["texto"], fonte)
            _fonte_na_marca(par, fonte)
            continue

        corrida = None  # qualquer outro bloco interrompe a numeração

        if tipo == "tabela":
            _tabela(doc, bloco["linhas"], fonte)
            continue

        texto = bloco["texto"]
        puro = _texto_puro(texto)

        if tipo == "titulo" and bloco["nivel"] <= 2 or puro in TITULOS_CENTRALIZADOS:
            _titulo(doc, texto, fonte)
        elif puro in DESTAQUE_VERMELHO:
            par = doc.add_paragraph()
            _escrever(par, texto, fonte, negrito=True, vermelho=True)
        else:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            # títulos menores (###) viram parágrafo em negrito
            _escrever(par, texto, fonte, negrito=True if tipo == "titulo" else None)

    doc.save(destino)
    return destino
